#!/usr/bin/env python3
"""信源清洗筛选器（零 token）：去重 → 评分 → 筛长期源 → 自动入采集器与 Excel。
输入 data-research/source_candidates.json（expand_sources.py 产出）；
输出：Excel 新增「候选源池」「长期源清单」sheet + harvest_sources.json 自动扩容。"""
from __future__ import annotations
import json, re, sys
from pathlib import Path
from urllib.parse import urlsplit
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
CAND = ROOT / "data-research" / "source_candidates.json"
XLSX = ROOT / "data-research" / "信息源总表-全量版.xlsx"
HSRC = ROOT / "tools" / "harvest_sources.json"
HARVEST_CAP = 70          # 采集器源总量上限（保证 15 分钟内跑完一轮）
AUTO_ADD_LIMIT = 36       # 本次自动并入采集器的长期源上限
LONG_TERM_TARGET = 120    # 长期源清单目标条数

BAD_PATH = re.compile(r"(github\.com/[^/]+/?$|github\.com/topics|/search\?|utm_|/login|/signup|\.pdf$)", re.I)

def norm(u: str) -> str:
    s = urlsplit(u.strip())
    host = s.netloc.lower().removeprefix("www.")
    path = re.sub(r"[?&]utm_[^&]*", "", s.path).rstrip("/")
    return host + path

def domain(u: str) -> str:
    return urlsplit(u).netloc.lower().removeprefix("www.")

def main() -> int:
    data = json.loads(CAND.read_text(encoding="utf-8"))
    cands = data["candidates"]
    print(f"输入候选: {len(cands)}")

    # 1) 去重（规范化 URL）
    seen: dict[str, dict] = {}
    for c in cands:
        n = norm(c["url"])
        if not n or BAD_PATH.search(n):
            continue
        if n not in seen or c.get("score", 0) > seen[n].get("score", 0):
            c = dict(c); c["norm"] = n
            seen[n] = c
    uniq = list(seen.values())
    print(f"去重后唯一源: {len(uniq)}（清洗掉 {len(cands) - len(uniq)} 条重复/无效路径）")

    # 2) 长期源筛选：可达 + official/repo 优先 + 有标题
    long_term = sorted(
        [c for c in uniq if c.get("status") == 200 and c.get("kind") in ("official", "repo")
         and c.get("title") and c.get("score", 0) >= 5
         and c.get("verified") != "offtopic" and c.get("verified") != "dead"],
        key=lambda c: (-c["score"], c["domain"] if "domain" in c else domain(c["url"])))
    for c in long_term:
        c["domain"] = domain(c["url"])
    # 同域去重：一个域最多保留 3 个页面（平台通常 1-2 个盯防页足够）
    from collections import Counter
    dc = Counter()
    picked = []
    for c in long_term:
        if dc[c["domain"]] < 3:
            picked.append(c); dc[c["domain"]] += 1
    long_term = picked[:LONG_TERM_TARGET]
    print(f"长期源筛选: {len(long_term)} 条（official/repo、200 可达、同域≤3）")

    # 3) 自动并入采集器（上限控制，跳过已存在域）
    cfg = json.loads(HSRC.read_text(encoding="utf-8"))
    have_domains = {domain(s["url"]) for s in cfg["sources"]}
    added = 0
    for c in long_term:
        if added >= AUTO_ADD_LIMIT or len(cfg["sources"]) >= HARVEST_CAP:
            break
        d = c["domain"]
        if d in have_domains:
            continue
        if "github.com" in d:
            pid = "gh-" + re.sub(r"[^a-z0-9]+", "-", c["url"].split("github.com/")[-1].lower())[:28]
            pats = ["free", "免费", "额度", "pricing"]
        else:
            pid = d.split(".")[0][:14] + "-" + (re.sub(r"[^a-z0-9]+", "", urlsplit(c["url"]).path)[:14] or "main")
            pats = ["免费", "赠送", "额度", "free", "pricing"]
        cfg["sources"].append({"id": pid[:30], "name": (c.get("title") or d)[:24],
                               "url": c["url"], "patterns": pats,
                               "note": f"长期源自动并入(分{c['score']}) 2026-10-28"})
        have_domains.add(d); added += 1
    HSRC.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"采集器并入 {added} 源，现共 {len(cfg['sources'])} 源")

    # 4) Excel：候选源池 + 长期源清单
    wb = load_workbook(XLSX)
    for name in ("候选源池", "长期源清单"):
        if name in wb.sheetnames:
            del wb[name]
    from openpyxl.styles import Font, PatternFill
    H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F6F4A")
    ws = wb.create_sheet("候选源池")
    ws.append(["url", "类型", "状态", "评分", "标题", "来源"])
    for c in sorted(uniq, key=lambda x: -x.get("score", 0)):
        ws.append([c["url"], c.get("kind"), c.get("status"), c.get("score"), c.get("title", ""), c.get("origin", "")[:40]])
    for j, w in enumerate([50, 10, 7, 7, 34, 26], 1): ws.column_dimensions[chr(64 + j)].width = w
    for cell in ws[1]: cell.font = H; cell.fill = HF
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("长期源清单")
    ws2.append(["url", "类型", "评分", "标题", "域名", "是否已入采集器"])
    added_urls = {s["url"] for s in cfg["sources"]}
    for c in long_term:
        ws2.append([c["url"], c.get("kind"), c.get("score"), c.get("title", ""), c["domain"],
                    "✅" if c["url"] in added_urls else "—"])
    for j, w in enumerate([50, 10, 7, 34, 24, 12], 1): ws2.column_dimensions[chr(64 + j)].width = w
    for cell in ws2[1]: cell.font = H; cell.fill = HF
    ws2.freeze_panes = "A2"
    log = wb["发现日志"]
    log.append(["2026-10-28", f"信源千人计划：扩展器产 {len(cands)} 候选→去重 {len(uniq)}→长期源 {len(long_term)}→采集器自动并入 {added}",
                "零token脚本(源里挖源)", "候选源池/长期源清单 sheet"])
    wb.save(XLSX)
    print(f"Excel 已更新: 候选源池 {len(uniq)} 行 / 长期源清单 {len(long_term)} 行")
    return 0

if __name__ == "__main__":
    sys.exit(main())
